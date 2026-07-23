#!/usr/bin/env python3
"""Send QC-filtered calls for one day to Telegram: summary + bad/uncertain only."""

from __future__ import annotations

import argparse
import os
import re
from datetime import datetime
from pathlib import Path
from typing import List, Tuple
from zoneinfo import ZoneInfo

from call_qc import CallAssessment, assess_call, format_day_summary
from docx import Document
from mango_sync import load_index
from mango_vpbx import MangoVpbxClient
from site_config import get_site, load_sites
from telegram_format import build_call_message, format_date_no_year, format_time_short
from telegram_notify import site_chat_ids, tg_broadcast_document, tg_broadcast_message
from site_lines import branch_label, is_tracked_call
from transcript_utils import parse_transcript_file, transcript_paragraphs


def load_dotenv(path: Path) -> None:
    if not path.exists():
        return
    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        key = key.strip()
        value = value.strip().strip('"').strip("'")
        if key and key not in os.environ:
            os.environ[key] = value


def call_direction(from_number: str, to_number: str) -> str:
    if "sip:" in (to_number or "") and "sip:" not in (from_number or ""):
        return "входящий"
    if "sip:" in (from_number or "") and "sip:" not in (to_number or ""):
        return "исходящий"
    return "неизвестно"


def write_docx(
    out_path: Path,
    date_str: str,
    time_str: str,
    direction: str,
    phone: str,
    transcript: List[Tuple[str, str]],
    branch: str = "",
) -> None:
    doc = Document()
    doc.add_heading(f"{format_date_no_year(date_str)} · {format_time_short(time_str)}", level=1)
    doc.add_paragraph(f"{direction} · {phone}")
    if branch:
        doc.add_paragraph(f"Точка: {branch}")
    doc.add_paragraph("")
    doc.add_heading("Расшифровка", level=2)
    for line in transcript_paragraphs(transcript):
        doc.add_paragraph(line)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--date", required=True, help="YYYY-MM-DD")
    parser.add_argument("--base-dir", default=".")
    parser.add_argument("--dotenv", default=".env")
    parser.add_argument("--send-all", action="store_true", help="Skip QC filter (debug)")
    parser.add_argument("--site", default="moscow", help="Site id: moscow, krasnoyarsk")
    args = parser.parse_args()

    base = Path(args.base_dir).resolve()
    load_dotenv(base / args.dotenv)
    token = os.getenv("TELEGRAM_BOT_TOKEN", "").strip()
    site = get_site(args.site)
    chat_ids = site_chat_ids(site)
    if not token or not chat_ids:
        raise SystemExit("Set TELEGRAM_BOT_TOKEN and TELEGRAM_CHAT_ID")

    target = datetime.strptime(args.date, "%Y-%m-%d").date()
    prefix = target.strftime("%Y-%m-%d")
    calls_dir = site.calls_dir(base)
    html_files = sorted(calls_dir.glob(f"{prefix}__*.html"))
    docx_dir = site.docx_dir(base)
    docx_dir.mkdir(parents=True, exist_ok=True)

    mango = MangoVpbxClient(site.api_key, site.api_salt)
    tz = ZoneInfo(site.tz_name)
    start = datetime(target.year, target.month, target.day, 0, 0, 0, tzinfo=tz)
    stats_unavailable = False
    try:
        calls = mango.fetch_stats(int(start.timestamp()), int(start.timestamp() + 86400 - 1))
        day_calls = [c for c in calls if datetime.fromtimestamp(c.start, tz).date() == target]
        tracked_calls = [c for c in day_calls if is_tracked_call(site.site_id, c)]
    except Exception as exc:
        print(f"[WARN] Mango stats: {exc}")
        stats_unavailable = True
        tracked_calls = []

    incoming = sum(
        1 for c in tracked_calls if call_direction(c.from_number, c.to_number) == "входящий"
    )
    outgoing = sum(
        1 for c in tracked_calls if call_direction(c.from_number, c.to_number) == "исходящий"
    )

    recording_calls = sorted(
        [c for c in tracked_calls if c.recording_ids],
        key=lambda c: c.start,
    )
    index = load_index(site.calls_index(base))
    call_pairs: List[Tuple[object, Path]] = []
    for call in recording_calls:
        meta = index.get(call.entry_id, {})
        fname = meta.get("file", "")
        html_path = calls_dir / fname if fname else None
        if not html_path or not html_path.exists():
            print(f"[WARN] no html for entry {call.entry_id[:12]}… ext={call.to_extension}")
            continue
        if not parse_transcript_file(html_path):
            continue
        call_pairs.append((call, html_path))

    analyzed = len(call_pairs)
    to_send: List[Tuple[object, Path, CallAssessment, List[Tuple[str, str]]]] = []

    for call, html_path in call_pairs:
        transcript = parse_transcript_file(html_path)
        if not transcript:
            continue
        assessment = assess_call(transcript)
        branch = branch_label(site.site_id, call)
        if args.send_all or assessment.should_send:
            to_send.append((call, html_path, assessment, transcript))
        else:
            print(f"[SKIP good] {branch} {html_path.name} — {assessment.category}")

    sent_bad = sum(1 for _, _, a, _ in to_send if a.verdict == "bad")
    sent_uncertain = sum(1 for _, _, a, _ in to_send if a.verdict == "uncertain")

    summary = format_day_summary(
        prefix,
        incoming,
        outgoing,
        analyzed,
        sent_bad,
        sent_uncertain,
        site_label=site.label,
        stats_unavailable=stats_unavailable,
    )
    tg_broadcast_message(
        token, chat_ids, summary, parse_mode="HTML" if stats_unavailable else None
    )

    if not to_send:
        tg_broadcast_message(token, chat_ids, f"{site.label}\n\nКосячных звонков за день не найдено.")
        print("[DONE] no problematic calls")
        return

    sent = 0
    total = len(to_send)
    for idx, (call, html_path, assessment, transcript) in enumerate(to_send, start=1):
        m = re.match(r"\d{4}-\d{2}-\d{2}__(\d{2}-\d{2}-\d{2})__(\d+)__", html_path.name)
        time_key = m.group(1) if m else ""
        phone = m.group(2) if m else "unknown"
        direction = call_direction(call.from_number, call.to_number)
        dt = datetime.fromtimestamp(call.start, tz)
        time_hms = dt.strftime("%H:%M:%S")
        branch = branch_label(site.site_id, call)

        comment = build_call_message(
            assessment,
            date_str=prefix,
            time_hms=time_hms,
            direction=direction,
            phone=phone,
            site_label=site.label,
            branch=branch,
        )
        docx_path = docx_dir / f"{prefix}_{time_key}_{phone}.docx"
        write_docx(docx_path, prefix, time_hms, direction, phone, transcript, branch=branch)

        tg_broadcast_message(token, chat_ids, comment)
        tg_broadcast_document(token, chat_ids, docx_path, site.label)
        sent += 1
        print(f"[OK] {assessment.verdict} {html_path.name}")

    print(f"[DONE] sent {sent} DOCX files (bad={sent_bad}, uncertain={sent_uncertain})")


if __name__ == "__main__":
    main()
