"""Parse bank statement CSV/Excel. Uses the bank's own category column."""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

DATE_HINTS = ("дата операции", "дата платежа", "дата", "date", "operation date", "booking")
AMOUNT_HINTS = (
    "сумма операции",
    "сумма платежа",
    "сумма в валюте счета",
    "сумма",
    "amount",
    "операция сумма",
)
DEBIT_HINTS = ("дебет", "списание", "расход", "debit", "withdrawal")
CREDIT_HINTS = ("кредит", "зачисление", "приход", "credit", "deposit")
CATEGORY_HINTS = ("категория", "статья", "category", "рубрика")
DESC_HINTS = (
    "описание",
    "назначение",
    "комментарий",
    "description",
    "merchant",
    "контрагент",
    "получатель",
)
MCC_HINTS = ("mcc",)
CARD_HINTS = ("номер карты", "карта", "счет", "счёт", "card", "account")
STATUS_HINTS = ("статус", "status", "состояние")
CURRENCY_HINTS = ("валюта операции", "валюта платежа", "валюта", "currency")

SKIP_STATUSES = {
    "failed",
    "declined",
    "canceled",
    "cancelled",
    "rejected",
    "ошибка",
    "отклонен",
    "отклонён",
    "отменена",
    "отменен",
    "отменён",
}

TRANSFER_CATEGORY_MARKERS = (
    "перевод",
    "переводы",
    "между своими",
    "c2c",
    "p2p",
    "сбп",
    "internal",
    "transfer",
)
EMPTY_CATEGORY_MARKERS = ("", "другое", "прочее", "без категории", "n/a", "na", "none", "-")
OWN_ACCOUNT_MARKERS = (
    "между своими",
    "на свою",
    "свой счет",
    "свой счёт",
    "own account",
    "копилка",
    "внутрибанковский перевод между счетами",
)

_DATE_FORMATS = (
    "%d.%m.%Y %H:%M:%S",
    "%d.%m.%Y %H:%M",
    "%d.%m.%Y",
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%d %H:%M",
    "%Y-%m-%d",
    "%d/%m/%Y %H:%M:%S",
    "%d/%m/%Y",
    "%m/%d/%Y",
)


@dataclass
class ParsedTx:
    posted_at: datetime
    amount: float
    currency: str = "RUB"
    category: str = ""
    description: str = ""
    mcc: str = ""
    card: str = ""
    status: str = ""
    needs_review: bool = False
    suggested_kind: str = "expense"  # expense | income | transfer
    suggested_internal: bool = False
    extra: dict = field(default_factory=dict)

    @property
    def posted_date(self) -> date:
        return self.posted_at.date()


class ParseError(ValueError):
    pass


def sniff_kind(path: str | Path, filename: str | None = None) -> str:
    """What the file actually is — iPhone often drops the .pdf extension."""
    path = Path(path)
    name = (filename or path.name).lower()
    suffix = Path(name).suffix.lower()
    head = b""
    try:
        head = path.read_bytes()[:16]
    except OSError:
        pass
    if head.startswith(b"\xff\xd8") or head.startswith(b"\x89PNG") or head.startswith(b"GIF8"):
        raise ParseError("Это фото или картинка. Нужен файл выписки: PDF, CSV или Excel.")
    if head.startswith(b"%PDF") or suffix == ".pdf":
        return "pdf"
    ole = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
    if head.startswith(ole) or suffix == ".xls":
        return "xls"
    if head.startswith(b"PK") or suffix == ".xlsx":
        return "xlsx"
    if suffix in {".csv", ".txt"}:
        return "csv"
    if b"\x00" in head[:16]:
        raise ParseError("Не похоже на банковскую выписку. Нужен PDF, CSV или Excel.")
    return "csv"


def parse_statement(path: str | Path, *, filename: str | None = None) -> list[ParsedTx]:
    path = Path(path)
    kind = sniff_kind(path, filename)
    if kind == "pdf":
        from .parse_pdf import parse_pdf_statement

        return parse_pdf_statement(path, filename=filename or path.name)
    if kind in {"xlsx", "xls"}:
        rows = _read_xlsx_rows(path)
    else:
        rows = _read_csv_rows(path)
    return rows_to_transactions(rows)


def rows_to_transactions(rows: list[list[object]]) -> list[ParsedTx]:
    if not rows:
        raise ParseError("Файл пустой")
    header_idx = _find_header_row(rows)
    header = [_norm_header(c) for c in rows[header_idx]]
    mapping = _map_columns(header)
    if "date" not in mapping:
        raise ParseError("Не нашёл колонку с датой. Нужна выписка с датой и суммой.")
    if "amount" not in mapping and not ("debit" in mapping or "credit" in mapping):
        raise ParseError("Не нашёл колонку с суммой.")

    out: list[ParsedTx] = []
    for raw in rows[header_idx + 1 :]:
        if _row_empty(raw):
            continue
        tx = _row_to_tx(raw, mapping)
        if tx is None:
            continue
        out.append(tx)
    if not out:
        raise ParseError("В файле нет ни одной операции")
    return out


def classify_review(category: str, description: str, amount: float) -> tuple[bool, str, bool]:
    """Return (needs_review, suggested_kind, suggested_internal)."""
    cat = (category or "").strip().lower()
    desc = (description or "").strip().lower()
    blob = f"{cat} {desc}"

    is_own = any(m in blob for m in OWN_ACCOUNT_MARKERS)
    is_transfer = is_own or any(m in cat for m in TRANSFER_CATEGORY_MARKERS)
    if not is_transfer:
        is_transfer = "перевод" in desc or "сбп" in desc

    empty_cat = cat in EMPTY_CATEGORY_MARKERS

    if is_own:
        return False, "transfer", True
    if is_transfer or empty_cat:
        kind = "income" if amount > 0 else "expense"
        if is_transfer:
            kind = "transfer"
        return True, kind, False
    if amount > 0:
        return False, "income", False
    return False, "expense", False


def _read_csv_rows(path: Path) -> list[list[object]]:
    text = _decode_bytes(path.read_bytes())
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,|\t,")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";" if sample.count(";") >= sample.count(",") else ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    return [list(row) for row in reader]


def _read_xlsx_rows(path: Path) -> list[list[object]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows: list[list[object]] = []
        for row in ws.iter_rows(values_only=True):
            rows.append(list(row))
        return rows
    finally:
        wb.close()


def _decode_bytes(raw: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "cp1251"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _norm_header(cell: object) -> str:
    text = "" if cell is None else str(cell)
    text = text.replace("\ufeff", "").strip().lower()
    text = re.sub(r"\s+", " ", text)
    return text


def _find_header_row(rows: list[list[object]]) -> int:
    best_i, best_score = 0, -1
    for i, row in enumerate(rows[:40]):
        cells = [_norm_header(c) for c in row]
        if not any(cells):
            continue
        score = 0
        joined = " ".join(cells)
        if any(h in joined for h in DATE_HINTS):
            score += 3
        if any(h in joined for h in AMOUNT_HINTS):
            score += 3
        if any(h in joined for h in CATEGORY_HINTS):
            score += 2
        if any(h in joined for h in DESC_HINTS):
            score += 1
        if score > best_score:
            best_score = score
            best_i = i
    if best_score < 3:
        raise ParseError(
            "Не похоже на банковскую выписку: нет строки с датой и суммой."
        )
    return best_i


def _map_columns(header: list[str]) -> dict[str, int]:
    mapping: dict[str, int] = {}

    def pick(key: str, hints: Iterable[str]) -> None:
        if key in mapping:
            return
        used = set(mapping.values())
        exact: list[int] = []
        fuzzy: list[int] = []
        for i, name in enumerate(header):
            if i in used or not name:
                continue
            if name in hints:
                exact.append(i)
            elif any(name.startswith(h) or h in name for h in hints):
                fuzzy.append(i)
        if exact:
            mapping[key] = exact[0]
        elif fuzzy:
            mapping[key] = fuzzy[0]

    pick("date", DATE_HINTS)
    pick("amount", AMOUNT_HINTS)
    pick("debit", DEBIT_HINTS)
    pick("credit", CREDIT_HINTS)
    pick("category", CATEGORY_HINTS)
    pick("description", DESC_HINTS)
    pick("mcc", MCC_HINTS)
    pick("card", CARD_HINTS)
    pick("status", STATUS_HINTS)
    pick("currency", CURRENCY_HINTS)
    return mapping


def _row_empty(raw: list[object]) -> bool:
    return not any(str(c).strip() for c in raw if c is not None)


def _cell(raw: list[object], idx: int | None) -> object:
    if idx is None or idx < 0 or idx >= len(raw):
        return None
    return raw[idx]


def _row_to_tx(raw: list[object], mapping: dict[str, int]) -> ParsedTx | None:
    posted = parse_datetime(_cell(raw, mapping.get("date")))
    if posted is None:
        return None

    amount = parse_amount(_cell(raw, mapping.get("amount")))
    if amount is None:
        debit = parse_amount(_cell(raw, mapping.get("debit"))) or 0.0
        credit = parse_amount(_cell(raw, mapping.get("credit"))) or 0.0
        if debit and credit:
            amount = credit - debit
        elif credit:
            amount = abs(credit)
        elif debit:
            amount = -abs(debit)
        else:
            return None
    if abs(amount) < 0.0001:
        return None

    status = _as_str(_cell(raw, mapping.get("status")))
    if status.lower() in SKIP_STATUSES:
        return None

    category = _as_str(_cell(raw, mapping.get("category")))
    description = _as_str(_cell(raw, mapping.get("description")))
    needs_review, kind, internal = classify_review(category, description, amount)
    return ParsedTx(
        posted_at=posted,
        amount=amount,
        currency=_as_str(_cell(raw, mapping.get("currency"))) or "RUB",
        category=category,
        description=description,
        mcc=_as_str(_cell(raw, mapping.get("mcc"))),
        card=_as_str(_cell(raw, mapping.get("card"))),
        status=status,
        needs_review=needs_review,
        suggested_kind=kind,
        suggested_internal=internal,
    )


def parse_datetime(value: object) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        # Excel serial — skip; openpyxl usually gives datetime
        return None
    text = str(value).strip()
    if not text:
        return None
    text = text.replace("T", " ")
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text[: len(fmt) + 8], fmt)
        except ValueError:
            continue
    m = re.match(r"(\d{1,2})[.](\d{1,2})[.](\d{4})", text)
    if m:
        try:
            return datetime(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        except ValueError:
            return None
    return None


def parse_amount(value: object) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip()
    if not s:
        return None
    s = s.replace("\xa0", " ").replace(" ", "")
    s = re.sub(r"[₽€$]|руб\.?|RUB|RUR|USD|EUR", "", s, flags=re.I)
    s = s.strip()
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    if s.count(",") == 1 and s.count(".") == 0:
        s = s.replace(",", ".")
    elif s.count(",") == 1 and s.count(".") >= 1:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif s.count(",") > 1:
        s = s.replace(",", "")
    try:
        n = float(s)
    except ValueError:
        return None
    return -n if neg else n


def _as_str(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()
